"""네이버 검색 결과의 뉴스 기사를 크롤링하는 PyQt6 GUI 앱.

- 검색어를 입력하면 네이버 검색 결과에서 기사 제목 / 언론사 / 요약 / 링크를 수집
- '본문도 수집'을 체크하면 각 기사 페이지에서 본문까지 가져옴
- 목록에서 기사를 선택하면 오른쪽에 내용을 표시, 더블클릭하면 브라우저로 열기
- 결과를 Excel(.xlsx) 또는 JSON 파일로 저장

필요 패키지: pip install PyQt6 requests beautifulsoup4 openpyxl
"""
import json
import sys
import time
from urllib.parse import quote, urlparse

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from bs4 import BeautifulSoup
from PyQt6.QtCore import QThread, QUrl, Qt, pyqtSignal
from PyQt6.QtGui import QDesktopServices
from PyQt6.QtWidgets import (
    QApplication, QCheckBox, QFileDialog, QHBoxLayout, QHeaderView, QLabel,
    QLineEdit, QMainWindow, QMessageBox, QProgressBar, QPushButton, QSplitter,
    QTableWidget, QTableWidgetItem, QTextEdit, QVBoxLayout, QWidget,
    QAbstractItemView,
)

SEARCH_URL = ("https://search.naver.com/search.naver?where=nexearch&sm=top_hty"
              "&fbm=0&ie=utf8&query={query}")
HEADERS = {
    "User-Agent": ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                   "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"),
    "Accept-Language": "ko-KR,ko;q=0.9",
}
DELAY = 0.5  # 기사 요청 사이 대기(초): 서버 부담을 줄이기 위함
# 언론사별 본문 영역 후보. 못 찾으면 og:description(요약)으로 대체한다.
BODY_SELECTORS = ("#dic_area, #newsct_article, #articleBodyContents, #articleBody, "
                  "#articeBody, #article-view-content-div, .article_body, .article-body, "
                  '[itemprop="articleBody"]')


# ---------------------------------------------------------------- 크롤링 로직
def get_soup(url):
    res = requests.get(url, headers=HEADERS, timeout=10)
    res.raise_for_status()
    res.encoding = res.apparent_encoding if res.encoding == "ISO-8859-1" else res.encoding
    return BeautifulSoup(res.text, "html.parser")


def clean(text):
    return text.replace("새 창 열림", "").strip()


def parse_search(soup):
    """검색 결과 페이지에서 기사 목록을 추출한다.

    네이버는 class 이름이 해시값이라 자주 바뀌므로,
    변하지 않는 data-heatmap-target 속성(.tit, .body)을 기준으로 찾는다.
    """
    articles, seen = [], set()
    for a in soup.select('a[data-heatmap-target=".tit"]'):
        link = a.get("href")
        if not link or link in seen:
            continue
        seen.add(link)

        # 제목이 이 기사 하나뿐인 가장 큰 부모 블록을 찾는다.
        block = a.parent
        while block.parent and block.parent.name not in ("body", "html"):
            if len(block.parent.select('a[data-heatmap-target=".tit"]')) > 1:
                break
            block = block.parent

        body = block.select_one('a[data-heatmap-target=".body"]')
        press = block.select_one('[class*="profile-info-title-text"]')
        articles.append({
            "title": clean(a.get_text(strip=True)),
            "press": clean(press.get_text(strip=True)) if press else urlparse(link).netloc,
            "summary": body.get_text(strip=True) if body else "",
            "url": link,
            "content": "",
        })
    return articles


def fetch_body(url):
    """기사 페이지에서 본문을 가져온다."""
    soup = get_soup(url)
    node = soup.select_one(BODY_SELECTORS)
    if node:
        for tag in node.select("script, style, .img_desc, figure"):
            tag.decompose()
        return node.get_text("\n", strip=True)
    meta = soup.select_one('meta[property="og:description"]')
    return meta["content"].strip() if meta and meta.get("content") else ""


def save_excel(path, articles):
    """기사 목록을 엑셀 파일로 저장한다."""
    wb = Workbook()
    ws = wb.active
    ws.title = "뉴스"
    headers = ["번호", "제목", "언론사", "요약", "본문", "URL"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    def safe(text):
        # 엑셀이 허용하지 않는 제어문자 제거, 셀 글자수 한도(32767) 적용
        text = "".join(ch for ch in text if ch in "\n\t" or ord(ch) >= 32)
        return text[:32000]

    for i, art in enumerate(articles, 1):
        ws.append([i, safe(art["title"]), safe(art["press"]), safe(art["summary"]),
                   safe(art["content"]), art["url"]])
        ws.cell(row=i + 1, column=6).hyperlink = art["url"]
        ws.cell(row=i + 1, column=6).font = Font(color="0563C1", underline="single")

    for col, width in zip("ABCDEF", (6, 50, 16, 60, 80, 50)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (2, 4, 5))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


# ------------------------------------------------------------ 백그라운드 작업
class CrawlWorker(QThread):
    """GUI가 멈추지 않도록 크롤링을 별도 스레드에서 실행한다."""
    status = pyqtSignal(str)
    found = pyqtSignal(list)            # 검색 결과 목록
    article_done = pyqtSignal(int, str)  # (행 번호, 본문)
    progress = pyqtSignal(int, int)
    failed = pyqtSignal(str)

    def __init__(self, query, with_body):
        super().__init__()
        self.query = query
        self.with_body = with_body

    def run(self):
        try:
            self.status.emit("검색 중...")
            soup = get_soup(SEARCH_URL.format(query=quote(self.query)))
            articles = parse_search(soup)
            self.found.emit(articles)
            if not articles:
                self.status.emit("검색된 기사가 없습니다.")
                return
            if not self.with_body:
                self.status.emit(f"기사 {len(articles)}건 수집 완료")
                return
            total = len(articles)
            for i, art in enumerate(articles):
                if self.isInterruptionRequested():
                    self.status.emit("중단되었습니다.")
                    return
                self.status.emit(f"본문 수집 중... ({i + 1}/{total})")
                try:
                    content = fetch_body(art["url"])
                except requests.RequestException:
                    content = ""
                self.article_done.emit(i, content)
                self.progress.emit(i + 1, total)
                time.sleep(DELAY)
            self.status.emit(f"기사 {total}건 수집 완료")
        except requests.RequestException as e:
            self.failed.emit(f"요청 실패: {e}")


# ------------------------------------------------------------------- GUI
class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("네이버 뉴스 크롤러")
        self.resize(1100, 640)
        self.articles = []
        self.worker = None

        # 상단 검색줄
        self.query_edit = QLineEdit("반도체")
        self.query_edit.setPlaceholderText("검색어를 입력하세요")
        self.query_edit.returnPressed.connect(self.start_crawl)
        self.body_check = QCheckBox("본문도 수집")
        self.body_check.setChecked(True)
        self.search_btn = QPushButton("검색")
        self.search_btn.clicked.connect(self.start_crawl)
        self.stop_btn = QPushButton("중지")
        self.stop_btn.setEnabled(False)
        self.stop_btn.clicked.connect(self.stop_crawl)
        self.excel_btn = QPushButton("엑셀 저장")
        self.excel_btn.setEnabled(False)
        self.excel_btn.clicked.connect(self.save_excel_file)
        self.save_btn = QPushButton("JSON 저장")
        self.save_btn.setEnabled(False)
        self.save_btn.clicked.connect(self.save_json)

        top = QHBoxLayout()
        top.addWidget(QLabel("검색어"))
        top.addWidget(self.query_edit, 1)
        top.addWidget(self.body_check)
        top.addWidget(self.search_btn)
        top.addWidget(self.stop_btn)
        top.addWidget(self.excel_btn)
        top.addWidget(self.save_btn)

        # 기사 목록
        self.table = QTableWidget(0, 3)
        self.table.setHorizontalHeaderLabels(["제목", "언론사", "본문"])
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.table.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        self.table.verticalHeader().setVisible(False)
        header = self.table.horizontalHeader()
        header.setSectionResizeMode(0, QHeaderView.ResizeMode.Stretch)
        header.setSectionResizeMode(1, QHeaderView.ResizeMode.ResizeToContents)
        header.setSectionResizeMode(2, QHeaderView.ResizeMode.ResizeToContents)
        self.table.itemSelectionChanged.connect(self.show_detail)
        self.table.cellDoubleClicked.connect(self.open_in_browser)

        # 기사 상세
        self.detail_title = QLabel("기사를 선택하세요")
        self.detail_title.setWordWrap(True)
        self.detail_title.setStyleSheet("font-size: 16px; font-weight: bold;")
        self.detail_url = QLabel()
        self.detail_url.setOpenExternalLinks(True)
        self.detail_url.setTextInteractionFlags(Qt.TextInteractionFlag.TextBrowserInteraction)
        self.detail_url.setWordWrap(True)
        self.detail_text = QTextEdit()
        self.detail_text.setReadOnly(True)
        right = QWidget()
        rl = QVBoxLayout(right)
        rl.setContentsMargins(0, 0, 0, 0)
        rl.addWidget(self.detail_title)
        rl.addWidget(self.detail_url)
        rl.addWidget(self.detail_text, 1)

        splitter = QSplitter()
        splitter.addWidget(self.table)
        splitter.addWidget(right)
        splitter.setSizes([560, 540])

        self.progress = QProgressBar()
        self.progress.setVisible(False)

        central = QWidget()
        layout = QVBoxLayout(central)
        layout.addLayout(top)
        layout.addWidget(splitter, 1)
        layout.addWidget(self.progress)
        self.setCentralWidget(central)
        self.statusBar().showMessage("검색어를 입력하고 [검색]을 누르세요.")

    # ---- 크롤링 제어
    def start_crawl(self):
        query = self.query_edit.text().strip()
        if not query:
            QMessageBox.warning(self, "알림", "검색어를 입력하세요.")
            return
        if self.worker and self.worker.isRunning():
            return
        self.articles = []
        self.table.setRowCount(0)
        self.detail_title.setText("기사를 선택하세요")
        self.detail_url.clear()
        self.detail_text.clear()
        self.save_btn.setEnabled(False)
        self.excel_btn.setEnabled(False)

        self.worker = CrawlWorker(query, self.body_check.isChecked())
        self.worker.status.connect(self.statusBar().showMessage)
        self.worker.found.connect(self.on_found)
        self.worker.article_done.connect(self.on_article_done)
        self.worker.progress.connect(self.on_progress)
        self.worker.failed.connect(self.on_failed)
        self.worker.finished.connect(self.on_finished)
        self.search_btn.setEnabled(False)
        self.stop_btn.setEnabled(True)
        self.progress.setVisible(False)
        self.worker.start()

    def stop_crawl(self):
        if self.worker:
            self.worker.requestInterruption()

    def on_found(self, articles):
        self.articles = articles
        self.table.setRowCount(len(articles))
        for row, art in enumerate(articles):
            self.table.setItem(row, 0, QTableWidgetItem(art["title"]))
            self.table.setItem(row, 1, QTableWidgetItem(art["press"]))
            self.table.setItem(row, 2, QTableWidgetItem("" if self.body_check.isChecked() else "-"))
        self.save_btn.setEnabled(bool(articles))
        self.excel_btn.setEnabled(bool(articles))
        if articles and self.body_check.isChecked():
            self.progress.setRange(0, len(articles))
            self.progress.setValue(0)
            self.progress.setVisible(True)

    def on_article_done(self, row, content):
        self.articles[row]["content"] = content
        self.table.setItem(row, 2, QTableWidgetItem("✔" if content else "✘"))
        if self.table.currentRow() == row:
            self.show_detail()

    def on_progress(self, done, total):
        self.progress.setValue(done)

    def on_failed(self, message):
        QMessageBox.critical(self, "오류", message)
        self.statusBar().showMessage(message)

    def on_finished(self):
        self.search_btn.setEnabled(True)
        self.stop_btn.setEnabled(False)
        self.progress.setVisible(False)

    # ---- 표시 / 저장
    def show_detail(self):
        row = self.table.currentRow()
        if not 0 <= row < len(self.articles):
            return
        art = self.articles[row]
        self.detail_title.setText(f"{art['title']}  ({art['press']})")
        self.detail_url.setText(f'<a href="{art["url"]}">{art["url"]}</a>')
        text = art["content"] or art["summary"] or "(내용 없음)"
        self.detail_text.setPlainText(text)

    def open_in_browser(self, row, _column):
        if 0 <= row < len(self.articles):
            QDesktopServices.openUrl(QUrl(self.articles[row]["url"]))

    def save_excel_file(self):
        path, _ = QFileDialog.getSaveFileName(self, "엑셀 저장", f"{self.query_edit.text().strip() or 'news'}_뉴스.xlsx",
                                              "Excel (*.xlsx)")
        if not path:
            return
        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"
        try:
            save_excel(path, self.articles)
        except PermissionError:
            QMessageBox.critical(self, "오류", "파일을 저장할 수 없습니다. 엑셀에서 열려 있다면 닫고 다시 시도하세요.")
            return
        except OSError as e:
            QMessageBox.critical(self, "오류", f"저장 실패: {e}")
            return
        self.statusBar().showMessage(f"엑셀 저장 완료: {path}")

    def save_json(self):
        path, _ = QFileDialog.getSaveFileName(self, "JSON 저장", "news_result.json",
                                              "JSON (*.json)")
        if not path:
            return
        with open(path, "w", encoding="utf-8") as f:
            json.dump(self.articles, f, ensure_ascii=False, indent=2)
        self.statusBar().showMessage(f"저장 완료: {path}")

    def closeEvent(self, event):
        if self.worker and self.worker.isRunning():
            self.worker.requestInterruption()
            self.worker.wait(3000)
        event.accept()


if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = MainWindow()
    window.show()
    sys.exit(app.exec())
